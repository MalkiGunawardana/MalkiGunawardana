import tkinter as tk
from tkinter import *
from tkinter import filedialog
from tkinter import messagebox
from PIL import Image, ImageTk
from tkinter import ttk
from tkinter.ttk import Combobox
from openpyxl import Workbook
import pickle


background="#09567C"
frameBg="black"
frameFg="black"

root =tk.Tk()
root.title("Student management System")
root.iconbitmap("F:\document\SUSL\Projects\Lab Projects\images\logo.ico")
root.resizable(False,False)
root.config(bg=background)

width = 1300
height = 700

screenWidth = root.winfo_screenwidth()
screenHeight = root.winfo_screenheight()
x = int((screenWidth / 2) - (width / 2))
y = int((screenHeight / 2) - (height / 2))
root.geometry(f"{width}x{height}+{(x)}+{(y)}")

#functions

def uploadPic():
    filePath = filedialog.askopenfilename()
    if filePath:
        img = Image.open(filePath)
        img = img.resize((120, 120), Image.ANTIALIAS)
        img = ImageTk.PhotoImage(img)
        labelUpldImg.config(image=img)
        labelUpldImg.image = img

def addData():
    
    data = [
        entryName.get(),
        entryRegNo.get(),
        entryCon.get(),
        gender.get(),
        entryYear.get(),
        entryEnDate.get(),
        
    ]

    gridView.insert('', tk.END, values=data)
    clearEntries()
    saveData()
    

def searchData():
    srchData = search.get().lower().strip()
    if srchData:
        found = False
        for data in gridView.get_children():
            values = gridView.item(data)['values']
            if srchData in str(values[0]).lower() or srchData in str(values[1]).lower():
                gridView.selection_set(data)
                gridView.focus(data)
                entryName.delete(0, tk.END)
                entryName.insert(0, values[0])
                entryRegNo.delete(0, tk.END)
                entryRegNo.insert(0, values[1])
                entryCon.delete(0, tk.END)
                entryCon.insert(0, values[2])
                gender.set(values[3])
                entryYear.delete(0, tk.END)
                entryYear.insert(0, values[4])
                entryEnDate.delete(0, tk.END)
                entryEnDate.insert(0, values[5])
                found = True
                break

        if not found:
            messagebox.showinfo("Warning", "No matching data found")
    else:
        messagebox.showinfo("Warning", "Please enter a search term")



def updateData():
    selectedItem = gridView.selection()
    if selectedItem:
        gridView.item(selectedItem, values=(
            entryName.get(),
            entryRegNo.get(),
            entryCon.get(),
            gender.get(),
            entryYear.get(),
            entryEnDate.get()
        ))
        clearEntries()
        saveData()

def deleteData():
    selectedItem = gridView.selection()
    if selectedItem:
        gridView.delete(selectedItem) 
        clearEntries() 
        saveData()

def clearEntries():
    entryName.delete(0, tk.END)
    entryRegNo.delete(0, tk.END)
    entryCon.delete(0, tk.END)
    gender.set("")
    entryYear.delete(0, tk.END)
    entryEnDate.delete(0, tk.END) 
    labelUpldImg.config(image=upImg)

def exitApp():
    root.destroy()    
 

#image
topImg = Image.open(r"F:\document\SUSL\Projects\Lab Projects\images\topimg.png")
resizeTopImg = topImg.resize((40,40))    
topImgLogo = ImageTk.PhotoImage(resizeTopImg)
image_label = Label(root, image=topImgLogo,bd=0)
image_label.pack(side=TOP)

#frame
text_frame = Frame(root)
text_frame.pack(side=TOP, fill=X)

labelApex = Label(text_frame, text="APEX INSTITUTE", width=20, bg="white", fg='black', font='arial 25 bold')
labelApex.pack(side=TOP, fill=X)

labelReg = Label(text_frame, text="Student Registration", width=15, bg="white", fg='blue', font='arial 16')
labelReg.pack(side=TOP, fill=X)

#search box
search = StringVar()
entrySrch=Entry(root,textvariable=search,width=15,bd=0,font='arial 14')
entrySrch.place(x=1120,y=10)

#search image/button
searchImg = Image.open(r"F:\document\SUSL\Projects\Lab Projects\images\srchImg.png")
resizeSrchImg = searchImg.resize((20, 20))
imgSrch=ImageTk.PhotoImage(resizeSrchImg)
btnSrch=Button(root,compound=LEFT,image=imgSrch,font="arial 13 bold",bd=0,command=searchData)
btnSrch.pack()
btnSrch.place(x=1092,y=12)

#frame for border
frameBorder = tk.Frame(root, bg='white')
frameBorder.pack(side=LEFT, fill=Y, padx=4, pady=4)
frameBorder.config(width=35)

#frame2
frame2 = Frame(frameBorder,bg='white',highlightthickness=2,highlightbackground="#09567C")
frame2.pack(side=LEFT, fill=Y,)
frame2.config(width=35)

#upload image
uploadImg = Image.open(r"F:\document\SUSL\Projects\Lab Projects\images\uploadimg.png")
resizeUpImg = uploadImg.resize((120, 120))
upImg = ImageTk.PhotoImage(resizeUpImg)
labelUpldImg = Label(frame2,bg='black',image=upImg)
labelUpldImg.grid(row=0, column=0,pady=10,columnspan=2)


#bottom image
bottomImg = Image.open(r"F:\document\SUSL\Projects\Lab Projects\images\2.png")
resizeBottomImg = bottomImg.resize((500, 240))
BottImg = ImageTk.PhotoImage(resizeBottomImg)
labelBottImg = Label(root,bg='black',image=BottImg,bd=0)
labelBottImg.pack(side=BOTTOM)


#exit image/button
#exitImg = Image.open(r"F:\document\SUSL\Projects\Lab Projects\images\logout.png")  
#ResizeExitImg = exitImg.resize((30, 30))
#exit= ImageTk.PhotoImage(ResizeExitImg)
#exitButton = Button(root, image=exit, command=exitApp,bd=0,background="#09567C")
#exitButton.pack(side=BOTTOM)



#inputs

labelName = Label(frame2, text="Name", bg='white', fg='black', font='arial 14')
labelName.grid(row=1, column=0, padx=5, pady=5, sticky='w')

labelRegNo = Label(frame2, text="Registration No", bg='white', fg='black', font='arial 14')
labelRegNo.grid(row=2, column=0, padx=5, pady=5, sticky='w')

labelCon = Label(frame2, text="Contact No", bg='white', fg='black', font='arial 14')
labelCon.grid(row=3, column=0, padx=5, pady=5, sticky='w')

labelGender = Label(frame2, text="Gender", bg='white', fg='black', font='arial 14')
labelGender.grid(row=4, column=0, padx=5, pady=5, sticky='w')

labelYear = Label(frame2, text="Year", bg='white', fg='black', font='arial 14')
labelYear.grid(row=5, column=0, padx=5, pady=5, sticky='w')

labelEnDate = Label(frame2, text="En-Date", bg='white', fg='black', font='arial 14')
labelEnDate.grid(row=6, column=0, padx=5, pady=5, sticky='w')

regNo = StringVar()
date = StringVar()
name = StringVar()
conNo = StringVar()
gender = StringVar()
year = StringVar()
enDate = StringVar()

entryName = ttk.Entry(frame2, textvariable=name, font='arial 10')
entryName.grid(row=1, column=1, padx=10, pady=10,sticky='w')
entryName.config(width=25)

entryRegNo = ttk.Entry(frame2, textvariable=regNo, font='arial 10')
entryRegNo.grid(row=2, column=1, padx=10, pady=10,sticky='w')
entryRegNo.config(width=25)

entryCon = ttk.Entry(frame2, textvariable=conNo, font='arial 10')
entryCon.grid(row=3, column=1, padx=10, pady=10,sticky='w')
entryCon.config(width=25)

genderFrame = tk.Frame(frame2, bg='white')
genderFrame.grid(row=4, column=1, padx=10, pady=10,sticky='w')
genderFrame.config(width=25)

maleRadio = ttk.Radiobutton(genderFrame, text="Male", variable=gender, value="Male")
maleRadio.grid(row=0, column=0)

femaleRadio = ttk.Radiobutton(genderFrame, text="Female", variable=gender, value="Female")
femaleRadio.grid(row=0, column=1)

entryYear = ttk.Combobox(frame2, textvariable=year, font='arial 10 bold', values=["Year 1", "Year 2", "Year 3", "Year 4"])
entryYear.grid(row=5, column=1, padx=10, pady=10,sticky='w')
entryYear.config(width=22)
entryYear.set("Select Year")

entryEnDate = ttk.Entry(frame2, textvariable=enDate, width=25, font='arial 10')
entryEnDate.grid(row=6, column=1, padx=10, pady=10,sticky='w')
entryEnDate.config(width=25)

#buttons

style = ttk.Style()

style.configure('Custom.TButton', 
                font=('Arial bold', 10), 
                foreground='black', 
                background='#09567C',
                width=30, 
                padding=0)

btnUpload = ttk.Button(frame2, text="Upload Image", command=uploadPic,style='Custom.TButton')
btnUpload.grid(row=11, column=0, pady=(20,7), sticky='n',columnspan=2)

btnAdd = ttk.Button(frame2, text="Add", command=addData,style='Custom.TButton')
btnAdd.grid(row=12, column=0, pady=7, sticky='n',columnspan=2)

#btnSrch = ttk.Button(frame2, text="Search", command=searchData)
#btnSrch.grid(row=10, column=1, pady=10,padx=5, sticky='w')

btnUpdate = ttk.Button(frame2, text="Update", command=updateData,style='Custom.TButton')
btnUpdate.grid(row=13, column=0, pady=7,sticky='n',columnspan=2)

btnDel = ttk.Button(frame2, text="Delete", command=deleteData,style='Custom.TButton')
btnDel.grid(row=14, column=0, pady=7,sticky='n',columnspan=2)

#btnExit = ttk.Button(root, text="Exit", command=deleteData,style='Custom.TButton')
#btnExit.grid(row=14, column=0, pady=7,sticky='n',columnspan=2)



# Grid view to display data
columns = ('Name', 'Reg No', 'Contact No', 'Gender', 'Year', 'En-Date')
gridView = ttk.Treeview(root, columns=columns, show='headings')

for col in columns:
    gridView.heading(col, text=col)
    gridView.column(col, width=100)

gridView.pack(padx=4,pady=(4,10), fill='both', expand=True)



#store data
def saveData():
    data_to_save = [gridView.item(item)['values'] for item in gridView.get_children()]
    with open('data.pkl', 'wb') as file:
        pickle.dump(data_to_save, file)


def loadData():
    try:
        with open('data.pkl', 'rb') as file:
            data = pickle.load(file)
            for item in data:
                gridView.insert("", "end", values=item)
    except FileNotFoundError:
        pass

# Load data when the application starts
loadData()

root.mainloop()
